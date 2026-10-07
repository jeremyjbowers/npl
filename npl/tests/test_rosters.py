"""
Unit tests for npl.rosters. Builds a miniature roster workbook in memory so
the tests don't depend on the live Google Sheet.

Run with:  python manage.py test npl.tests.test_rosters
"""

import datetime
import io
import unittest

import openpyxl
from openpyxl.styles import Font, PatternFill

from npl import rosters


def fill(argb):
    return PatternFill(fill_type="solid", fgColor=argb)


def build_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Key")
    wb.create_sheet("2026 Contracts")

    ws = wb.create_sheet("Testers #99")
    ws["A1"] = "Testville Testers"
    ws["E1"] = "AMERICAN LEAGUE\nOIL CAN DIVISION"
    ws["K1"] = "f. 2009"
    ws["O3"] = "General Manager: Pat Example (pat@example.com)"
    ws["E4"], ws["G4"], ws["J4"], ws["K4"], ws["L4"], ws["M4"], ws["N4"] = (
        2026, 182462308, 0, 3315526, 3315526, 83, 40,
    )
    ws["E5"], ws["G5"], ws["J5"], ws["K5"], ws["N5"] = 2027, 140811621, 41650688, 8065526, 33

    header = ["", "", "", "", "POS", "MLB", "MLS", "OPT", "STA"] + list(range(2026, 2034)) + ["BUYOUT", "CONTRACT TERMS"]
    ws.append(header)  # row 6

    def row(values):
        ws.append(values)
        return ws.max_row

    row(["", "ACTIVE ROSTER"])
    row(["", "STARTING PITCHERS", "SS#", "MLB#"])
    r = row([1, "Holmes, Clay", 1018, 605280, "RHP", "NYM", 4.091, 3, None,
             6210000, 8280000, 10350000, 10350000, None, None, None, None, None, "Extension"])
    ws.cell(r, 10).font = Font(italic=True)  # 2026 covered by another team
    r = row([1, "Crawford, J.P.", 560, 641487, "SS", "SEA", 6.157, 99, "QO",
             8780000, 14048000, 13696800, None, None, None, None, None, 2739400, "vests with 450 ABs in 2027"])
    ws.cell(r, 12).fill = fill("FFFF0000")  # vesting option
    r = row([1, "Soler, Jorge", 659, 624585, "OF", "LAA", 10.028, 99, None,
             5767608, 7209500, None, None, None, None, None, None, 1081400, "Offseason FA (2026)", "extra note"])
    ws.cell(r, 11).fill = fill("FFFFFF01")  # near-yellow -> club option
    ws.cell(r, 2).font = Font(italic=True)  # signed as offseason FA
    r = row([1, "Ginn, J.T.", 60, 669372, "RHP", "ATH", 0.105, 2, None, 777000])
    ws.cell(r, 10).fill = fill("FFD9EAD3")  # pre-arb tier 0
    r = row([1, "Larnach, Trevor", 700, 663993, "OF", "MIN", 3.144, 1, None, 2216034.75, 52866300, 46237600])
    ws.cell(r, 10).fill = fill("FFCFE2F3")  # light blue 3 -> arbitration
    ws.cell(r, 11).fill = fill("FFD9D2E9")  # light purple 3 -> player option
    ws.cell(r, 12).fill = fill("FFEAD1DC")  # light magenta 3 -> player option / opt-out
    ws.cell(r, 7).fill = fill("FFB7B7B7")  # grey MLS cell -> pending contract terms
    row([1, "Broken, Row", "-", 669203, "RHP", "ARI", 6.082, 99, None, 11000000])
    row([1, "Shifted, Row", "-", "$669,203", "RHP"])  # salary slid into the MLB# column

    row(["", "TRIPLE-A UNIVERSAL BASEBALL ASSOCIATION", None, None, None, None, None, None, None,
         None, None, None, None, "RECALL DATE"])
    row(["", "ON OPTION"])
    r = row([1, "Foster, Cameron", 382, 671382, "RHP", "BAL", 0.0, 3, None, 775154,
             None, None, None, "8/31", None, None, None, None, "2026 NRI", "Bid sign"])
    row(["", "NON-ROSTER"])
    row(["-", "Robinson, Kristian", 5038, 677565, "OF", "ARI", 0.0, 1, None, 760000])
    row(["", "DOUBLE-A DAMON RUTHERFORD'S PIONEERS"])
    row(["-", "Haynes, Jagger", "-", 695247, "LHP", "SD", 2025])
    row(["", "Single A Timestream Legends"])
    row(["-", "Beam, Drew", "-", 701795, "RHP", "KC", "2028"])

    row(["", "FINANCIALS"])
    row(["", "TOTAL FINANCIALS"])
    row(["", "SALARY CAP", None, None, None, None, None, None, None, 182462308, 182462308])
    row(["", "$ Reserves Liabilities"])
    row(["", "Urias, Julio", None, None, "Release", None, None, None, None, 3080400])
    row(["", "Holmes, Clay", None, None, "Carried salary by Perwar", None, None, None, None, -6210000])
    row(["", "CASH INCOME + EXPENDITURES"])
    row(["", "Trade", None, None, "Testers > Bulldog", None, None, None, None, -75000])
    row(["", "IFA CAP SPACE", None, None, "Current season always assumed to be $4.75M (Base)."])
    row(["", None, None, None, "Base:", None, 4750000, None, None, "Can acquire:", 6762500, "Total:", 1550000])
    row(["", "Date", None, None, "Team To/From"])
    row(["", datetime.date(2025, 8, 4), None, None, "Testers > Bosses", None, None, None, None, -1700000])

    # A second team that lists Foster again, to exercise the duplicate warning.
    ws2 = wb.create_sheet("Others #98")
    ws2["A1"] = "Otherton Others"
    ws2.append([])
    ws2.append([])
    ws2.append([])
    ws2.append([])
    ws2.append(header)
    ws2.append(["", "ACTIVE ROSTER"])
    ws2.append([1, "Foster, Cameron", 382, 671382, "RHP", "BAL", 0.0, 3, None, 775154])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ParseWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.parsed = rosters.parse_workbook(build_workbook())
        cls.testers = cls.parsed["teams"][0]
        cls.by_id = {p["mlb_id"]: p for p in cls.testers["players"]}

    def test_only_team_tabs_are_parsed(self):
        self.assertEqual([t["tab"] for t in self.parsed["teams"]], ["Testers #99", "Others #98"])

    def test_team_header(self):
        team = self.testers["team"]
        self.assertEqual(team["name"], "Testville Testers")
        self.assertEqual(team["league"], "AMERICAN LEAGUE")
        self.assertEqual(team["division"], "OIL CAN DIVISION")
        self.assertEqual(team["founded"], 2009)
        self.assertEqual(team["staff"], [{"title": "General Manager", "name": "Pat Example", "email": "pat@example.com"}])
        self.assertEqual(team["seasons"]["2026"]["payroll"], 182462308)
        self.assertEqual(team["seasons"]["2027"]["cap_space"], 41650688)
        self.assertEqual(team["roster_counts"], {"85_man": 83, "40_man": 40, "30_man": 33})
        self.assertEqual(team["affiliates"]["AAA"], "UNIVERSAL BASEBALL ASSOCIATION")
        self.assertEqual(team["affiliates"]["A"], "Timestream Legends")

    def test_salary_years_come_from_header(self):
        self.assertEqual(self.testers["salary_years"], list(range(2026, 2034)))

    def test_mlb_player_basics(self):
        p = self.by_id[605280]
        self.assertEqual((p["first_name"], p["last_name"]), ("Clay", "Holmes"))
        self.assertEqual(p["scoresheet_id"], 1018)
        self.assertEqual(p["position"], "RHP")
        self.assertEqual(p["mlb_org"], "NYM")
        self.assertEqual(p["level"], "MLB")
        self.assertEqual(p["position_group"], "STARTING PITCHERS")
        self.assertTrue(p["on_40_man"])
        self.assertEqual(p["service_time"], {"raw": "4.091", "years": 4, "days": 91})
        self.assertIsNone(p["eligible_year"])
        self.assertEqual(p["options"], 3)
        self.assertEqual(p["contract_terms"], "Extension")
        self.assertEqual([s["year"] for s in p["salaries"]], [2026, 2027, 2028, 2029])
        self.assertEqual(p["salaries"][0]["amount"], 6210000)

    def test_covered_salary_is_flagged_by_italics(self):
        p = self.by_id[605280]
        self.assertTrue(p["salaries"][0]["covered"])
        self.assertFalse(p["salaries"][1]["covered"])

    def test_contract_type_from_fill(self):
        crawford = self.by_id[641487]
        self.assertEqual(crawford["status"], "QO")
        self.assertEqual(crawford["buyout"], 2739400)
        self.assertEqual([s["type"] for s in crawford["salaries"]], ["guaranteed", "guaranteed", "vesting_option"])
        self.assertEqual(self.by_id[669372]["salaries"][0]["type"], "pre_arb_0")

    def test_near_colour_snaps_to_legend(self):
        soler = self.by_id[624585]
        self.assertEqual(soler["salaries"][1]["type"], "club_option")
        self.assertEqual(soler["salaries"][1]["fill"], "FFFFFF01")
        self.assertTrue(soler["signed_as_offseason_fa"])
        self.assertEqual(soler["notes"], ["extra note"])

    def test_palette_variants_classified_by_hue(self):
        larnach = self.by_id[663993]
        self.assertEqual(
            [s["type"] for s in larnach["salaries"]],
            ["arbitration", "player_option", "player_option"],
        )
        self.assertEqual(larnach["salaries"][0]["amount"], 2216034.75)
        self.assertEqual(larnach["flags"], ["pending_contract_terms"])
        self.assertEqual(larnach["roster_flag"], "1")

    def test_aaa_recall_date_and_subsection(self):
        foster = self.by_id[671382]
        self.assertEqual(foster["level"], "AAA")
        self.assertEqual(foster["roster_subsection"], "on_option")
        self.assertEqual(foster["recall_date"], "8/31")
        self.assertEqual([s["year"] for s in foster["salaries"]], [2026])
        self.assertEqual(foster["contract_terms"], "2026 NRI")
        self.assertEqual(self.by_id[677565]["roster_subsection"], "non_roster")

    def test_minor_leaguers_carry_eligible_year(self):
        haynes = self.by_id[695247]
        self.assertEqual(haynes["level"], "AA")
        self.assertIsNone(haynes["service_time"])
        self.assertEqual(haynes["eligible_year"], 2025)
        self.assertFalse(haynes["on_40_man"])
        self.assertIsNone(haynes["scoresheet_id"])
        beam = self.by_id[701795]
        self.assertEqual(beam["level"], "A")
        self.assertEqual(beam["eligible_year"], 2028)

    def test_rows_without_mlb_id_are_skipped_with_warning(self):
        self.assertNotIn("Shifted, Row", [p["name"] for p in self.testers["players"]])
        self.assertTrue(any("Shifted, Row" in w for w in self.parsed["warnings"]))
        self.assertIn(669203, self.by_id)  # a '-' scoresheet id is fine

    def test_financials(self):
        fin = self.testers["financials"]
        self.assertEqual(fin["summary"]["SALARY CAP"], {"2026": 182462308, "2027": 182462308})
        self.assertEqual(
            fin["liabilities"],
            [
                {"player": "Urias, Julio", "description": "Release", "amount": 3080400},
                {"player": "Holmes, Clay", "description": "Carried salary by Perwar", "amount": -6210000},
            ],
        )
        self.assertEqual(fin["cash_ledger"], [{"label": "Trade", "description": "Testers > Bulldog", "amount": -75000}])
        self.assertEqual(
            fin["ifa"],
            {
                "base": 4750000,
                "can_acquire": 6762500,
                "total": 1550000,
                "trades": [{"date": "2025-08-04", "teams": "Testers > Bosses", "amount": -1700000}],
            },
        )

    def test_financial_rows_are_not_players(self):
        self.assertEqual(len(self.testers["players"]), 10)
        self.assertNotIn("Urias, Julio", [p["name"] for p in self.testers["players"]])

    def test_cross_team_duplicate_is_warned(self):
        self.assertTrue(any("671382" in w and "multiple teams" in w for w in self.parsed["warnings"]))

    def test_flatten_players(self):
        flat = rosters.flatten_players(self.parsed)
        self.assertEqual(len(flat), 11)
        self.assertEqual(flat[0]["npl_team"], "Testville Testers")


class HelperTests(unittest.TestCase):
    def test_to_money(self):
        self.assertEqual(rosters.to_money("$3,009,259"), 3009259)
        self.assertEqual(rosters.to_money("-$1,700,000"), -1700000)
        self.assertEqual(rosters.to_money(2474616.375), 2474616.375)
        self.assertEqual(rosters.to_money(140811620.5), 140811620.5)
        self.assertEqual(rosters.to_money(760000.0), 760000)
        self.assertIsInstance(rosters.to_money(760000.0), int)
        self.assertIsNone(rosters.to_money("Extension"))
        self.assertIsNone(rosters.to_money(None))

    def test_salary_type_for_fill(self):
        self.assertEqual(rosters.salary_type_for_fill(None), "guaranteed")
        self.assertEqual(rosters.salary_type_for_fill("FFFDFDFD"), "guaranteed")
        self.assertEqual(rosters.salary_type_for_fill("FFFFFF00"), "club_option")
        self.assertEqual(rosters.salary_type_for_fill("FFFF0000"), "vesting_option")
        self.assertEqual(rosters.salary_type_for_fill("FFF4CCCC"), "vesting_option")
        self.assertEqual(rosters.salary_type_for_fill("FFB4A7D6"), "player_option")
        self.assertEqual(rosters.salary_type_for_fill("FFA4C2F4"), "arbitration")
        self.assertEqual(rosters.salary_type_for_fill("FFC9DAF8"), "arbitration")
        self.assertEqual(rosters.salary_type_for_fill("FFB6D7A8"), "pre_arb_1")
        self.assertEqual(rosters.salary_type_for_fill("FF93C47D"), "pre_arb_2")

    def test_parse_service_time(self):
        self.assertEqual(rosters.parse_service_time(5.058), ({"raw": "5.058", "years": 5, "days": 58}, None))
        self.assertEqual(rosters.parse_service_time("8.17"), ({"raw": "8.170", "years": 8, "days": 170}, None))
        self.assertEqual(rosters.parse_service_time(0), ({"raw": "0.000", "years": 0, "days": 0}, None))
        self.assertEqual(rosters.parse_service_time(2029), (None, 2029))
        self.assertEqual(rosters.parse_service_time("2029"), (None, 2029))
        self.assertEqual(rosters.parse_service_time(None), (None, None))

    def test_to_mlb_id_rejects_scoresheet_ids_and_money(self):
        self.assertEqual(rosters.to_mlb_id(656288.0), 656288)
        self.assertEqual(rosters.to_mlb_id(80537), 80537)
        self.assertIsNone(rosters.to_mlb_id(5002))
        self.assertIsNone(rosters.to_mlb_id("$669,203"))
        self.assertIsNone(rosters.to_mlb_id("-"))

    def test_split_name(self):
        self.assertEqual(rosters.split_name("Witt Jr., Bobby"), ("Bobby", "Witt Jr."))
        self.assertEqual(rosters.split_name("Kiner-Falefa, Isiah"), ("Isiah", "Kiner-Falefa"))


if __name__ == "__main__":
    unittest.main()
